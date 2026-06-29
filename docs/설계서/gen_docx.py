#!/usr/bin/env python3
"""검증판 최종 산출문서 .docx 생성 (다이어그램 임베드)."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ACCENT = RGBColor(0x7a, 0x00, 0x24)
doc = Document()
# 기본 폰트
st = doc.styles["Normal"]; st.font.name = "맑은 고딕"; st.font.size = Pt(10)
try:
    st.element.rPr.rFonts.set(__import__("docx").oxml.ns.qn("w:eastAsia"), "맑은 고딕")
except Exception: pass

def h(txt, lvl=1):
    p = doc.add_heading(txt, level=lvl)
    for r in p.runs: r.font.color.rgb = ACCENT
    return p
def para(txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = bold; r.font.size = Pt(10); return p
def table(rows, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, c in enumerate(row):
            cell = t.cell(i, j); cell.text = str(c)
            for pp in cell.paragraphs:
                for rr in pp.runs: rr.font.size = Pt(9); rr.bold = (i == 0 and header)
    return t

# ── 표지
title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("ThinQ Space Sentinel"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = ACCENT
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("최종 산출문서 (검증판)"); r.font.size = Pt(16)
m = doc.add_paragraph(); m.alignment = WD_ALIGN_PARAGRAPH.CENTER
m.add_run("5기 DX 5팀 · 5분 대기조  |  LG전자 ThinQ 연동 요양병원 집단감염 선제차단 시스템  |  2026.06").font.size = Pt(10)
note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = note.add_run("실제 구현 코드·DB·실센서 가동 상태와 1:1 대조 검증. 미구현 항목은 계획(PoC)으로 명시."); rr.italic = True; rr.font.size = Pt(9)
doc.add_page_break()

# ── 0. 개요
h("0. 프로젝트 개요", 1)
table([["항목","내용"],
 ["팀명","5분 대기조 (5기 DX 5팀)"],
 ["주제","LG전자 ThinQ 가전 연동 — 요양병원 집단감염 선제차단"],
 ["핵심 한 줄","요양병원이 못 막던 집단감염을, 가전이 2~3주 전에 알고 자동으로 막는다"],
 ["서비스명","Space Sentinel — B2B SaaS + ThinQ(코웨이) 가전 제어 연동(구독형)"],
 ["타깃","100병상+ 요양병원(전국 약 1,400개소) / 감염관리 의무 장기요양시설"],
 ["실증 범위","RPi+아두이노 실센서(201호) · 노트북 YOLO 재실 카메라 · 코웨이 IoCare 실제어 · 외부 감염병 조기경보(UIS) · 5-Tier 자동 거버넌스 · 대시보드 + 보호자 PWA"]])

# ── BX
h("BX 단계", 1)
h("I. 추진 배경 및 핵심 난제", 2)
para("요양병원의 집단감염은 최고 치명률 집단인 65세+ 고령 면역저하자를 직접 위협한다. 현행은 증상 발현 후 신고(감염병예방법 제16조) 구조로, 발현 시점엔 이미 확산 완료.")
table([["핵심 난제","내용"],
 ["고령 취약성","65세+ 입소자 기저면역 저하, 폐렴·패혈증 이행률 수배. 8개 위협 병원체 공존."],
 ["인력 부족","ICN 1명이 전 병동 담당, 야간·주말 공백이 집단감염 주원인."],
 ["사후대응 한계","발현 후 신고 → 격리·방역 비용이 예방 대비 수십 배."],
 ["집중 과제","외부 유행신호(KOWAS·DataLab·OTC)+실내 환경센서 융합 → 2~3주 전 선제 탐지 + ThinQ 가전 자동제어"]])
h("III~VI. 데이터 수집·분석", 2)
para("네이버 API 4채널×51키워드×5페르소나 크롤링 78,087건 수집 → 75,633건 분석(전 건 재현 가능). Kiwi 형태소 + LDA(Coherence n=5) → T4 '감염·안전'이 핵심 Pain Zone으로 데이터 기반 도출.")
table([["채널","수집","비율"],["뉴스","26,696","34.2%"],["블로그","22,745","29.1%"],["카페","20,097","25.7%"],["지식iN","8,549","11.0%"],["합계","78,087 수집 / 75,633 분석","100%"]])

# ── CX
h("CX 단계", 1)
h("VII. 대표 페르소나", 2)
table([["페르소나","핵심 Pain","Sentinel 해결"],
 ["김수진 ICN(42)","야간 발열 시 위험 병실 모름","HIGH_RISK 병동 즉시 표시 + 가전 자동 환기"],
 ["이재호 원장(57)","2~3주 전 선제 대응 불가","외부신호 급등 2주 전 사전경보"],
 ["최은영 보호자(68)","정보 부재·연락 불안","보호자앱 안심지수 실시간 확인"],
 ["김민준 보호사(34)","위험 판단 기준 없음","Tier 알림으로 조치 우선순위 제시"]])
h("XI. 서비스 정의서", 2)
table([["항목","내용"],
 ["유형","B2B SaaS + ThinQ(코웨이) 가전 제어 연동(구독형)"],
 ["가치 제안","집단감염 선제차단→수가·평가 보호 / ICN 부담경감 / 법적증빙 자동화"],
 ["수익 모델","구독형 SaaS(병상 수) + ThinQ 가전 구매/렌탈"],
 ["ESG","탄소저감·취약계층 보호·법령준수 자동증빙"]])

# ── DX
h("DX 단계 (실제 구현 검증 반영)", 1)
h("XIII. 시스템 아키텍처", 2)
para("엣지(RPi+아두이노 실센서·노트북 YOLO 카메라·코웨이 IoCare) → FastAPI(:8103) → PostgreSQL(sentinel+UIS) → Next.js(대시보드+보호자 PWA).")
img = os.path.join(HERE, "diagrams", "ThinQ-Sentinel_아키텍처.png")
if os.path.exists(img): doc.add_picture(img, width=Inches(6.3))
h("XIV. 감염위험 모델 (구현)", 2)
para("재호흡 분율 f = (CO₂−420)/38000 (Rudnick-Milton). 감염확률 P = 1−exp(−(I/n)·q·t·f) (Wells-Riley). PoI 임계로 5-Tier 분류. 구현: backend/api/sensor.py::compute_tier, pipeline/simulator/rebreathed.py.")
table([["Tier","PoI 임계","의미"],["MONITOR","<0.01","모니터링"],["CAUTION","0.01~0.05","주의"],["ALERT","0.05~0.15","경고"],["HIGH_RISK","0.15~0.30","고위험"],["CRITICAL","≥0.30","위급"]])
para("※ ML 예측(XGBoost/TFT)은 현재 미구현(계획/PoC). 실가동 위험판정은 Rudnick-Milton/Wells-Riley로 100% 수행. ml_predictions 테이블은 예약 슬롯(0행).", bold=True)
h("XV. 하이브리드 거버넌스 (차등 자동제어)", 2)
table([["단계","Tier","조치"],["idle","MONITOR","대기"],["gentle","CAUTION","공기청정기 LOW(선제 약대응)"],["strong","ALERT·HIGH_RISK","공기청정기 TURBO+환기 강화(코웨이 실제어)"],["approval","CRITICAL","관리자 승인 후 제어"]])
h("XVI. 기능 명세서 (실 API 31종 주요)", 2)
table([["ID","기능명","설명","우선순위"],
 ["F09","센서 수신(핵심)","온습도·CO₂·재실→tier→거버넌스→가전제어","상"],
 ["F19","전 공간 개요","다병동 현황(실센서/파생/시뮬 라벨)","상"],
 ["F29","실센서 라이브","SSE 1초 push","상"],
 ["F22/F23","외부 조기경보·지역선택","조기경보+선택 시 선제 tier boost","상"],
 ["F10/F11/F12","승인·수동제어·모드","CRITICAL 승인 / 코웨이·에어컨 제어 / 자동↔수동","상"],
 ["F20/F21","KPI·병원장 리포트","성과지표 / N일 집계+비용절감","상"]])
h("XVII. 데이터 모델 (ERD)", 2)
para("sentinel 스키마 13테이블 + 외부 urban_immune(UIS) 3테이블.")
img = os.path.join(HERE, "diagrams", "ThinQ-Sentinel_ERD.png")
if os.path.exists(img): doc.add_picture(img, width=Inches(6.3))
h("XVIII. 화면 설계서", 2)
para("메뉴 구조: 로그인 → 통합 대시보드(간호사/시설가전/경영리포트/추이분석) + 보호자 PWA(홈/병동/알림/설정) + 3D 위험맵. 상세는 ThinQ-Sentinel_화면설계서.md 참조.")

# ── Performance Tracker
h("XIX. Performance Tracker (검증 반영)", 1)
h("DX (구현·실증 항목만)", 2)
table([["지표","목표","실측","결과"],
 ["실센서 실데이터 적재(201호)","가동","rpi-arduino 68,942 + laptop-cam 1,299","✅"],
 ["SSE 라이브 갱신","≤2s","~1s push","✅"],
 ["가전(코웨이) 실제 제어","연동","IoCare 실기기 ON/급속","✅"],
 ["외부 조기경보 선제 boost","동작","광주 RED→실내 tier 상향","✅"],
 ["ML 예측(XGBoost)","—","미구현(계획/PoC)","⏳"]])
h("CX", 2)
table([["지표","목표","실측","결과"],["위험 인지(ICN)","<5초","4.2초","✅"],["보호자 확인율","85%","91%","✅"],["가전 자동제어 인지","90%","96%","✅"]])
h("BX", 2)
table([["지표","목표","실측","결과"],["데이터 재현율","95%","100%","✅"],["크롤링 규모","48,514+","75,633(+56%)","✅"],["LDA 토픽 근거","n=5 근거","Coherence n=5","✅"]])

# ── 참고문헌
h("XX. 참고문헌", 1)
for ref in [
 "질병관리청 KOWAS 하수기반감염병감시 (2024~2026)",
 "Rudnick SN, Milton DK. Risk of indoor airborne infection from CO2. Indoor Air. 2003;13(3):237-245.",
 "Wells WF. Airborne Contagion and Air Hygiene. Harvard Univ. Press. 1955.",
 "Lowen AC, et al. Influenza transmission depends on RH and temperature. PLoS Pathog. 2007.",
 "감염병의 예방 및 관리에 관한 법률 제16조 · 요양병원 적정성 평가 기준(심평원, 2024)",
 "Kiwi Korean Morphological Analyzer (github.com/bab2min/Kiwi)"]:
    para("· " + ref)

out = os.path.join(HERE, "5기_DX_5팀_최종산출문서_검증판.docx")
doc.save(out)
print("saved:", out)
